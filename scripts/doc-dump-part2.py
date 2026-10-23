#!/usr/bin/env python3
"""Dump full text of Product Design 2.docx with paragraph indices and heading colors."""
import sys
from docx import Document
from docx.shared import RGBColor

path = "/home/z/my-project/upload/Product Design 2.docx"
doc = Document(path)

out = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if not text:
        continue
    style = p.style.name if p.style else ""
    # detect red-ish color runs (Part headings were red in prior docs)
    color = ""
    for run in p.runs:
        c = run.font.color
        if c and c.rgb is not None:
            r, g, b = c.rgb[0], c.rgb[1], c.rgb[2]
            if r > 150 and g < 100 and b < 100:
                color = "RED"
            break
    out.append(f"[{i:04d}]{'[' + color + ']' if color else ''}({style}) {text}")

print(f"TOTAL non-empty paragraphs: {len(out)}")
print("=" * 80)
print("\n".join(out))

# tables if any
for ti, table in enumerate(doc.tables):
    print(f"\n===== TABLE {ti} =====")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        print(" | ".join(cells))
