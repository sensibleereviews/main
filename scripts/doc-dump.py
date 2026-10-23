"""Dump a .docx to indexed text: paragraphs (with red-heading + style tags) and tables."""
import sys
from docx import Document


def run_hex(run):
    try:
        c = run.font.color
        if c is not None and c.rgb is not None:
            return str(c.rgb)  # e.g. 'FF0000'
    except Exception:
        pass
    return None


def is_red(run):
    h = run_hex(run)
    if not h or len(h) != 6:
        return False
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return r > 140 and g < 110 and b < 110


def dump(path):
    doc = Document(path)
    lines = []
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        red = "[RED]" if any(is_red(r) for r in p.runs if r.text.strip()) else ""
        style = p.style.name if p.style is not None else ""
        st = f" <{style}>" if style and style != "Normal" else ""
        lines.append(f"{i:4d} {red}{st}: {t}")
    print("\n".join(lines))
    print(f"\n=== non-empty paragraphs: {len(lines)} | tables: {len(doc.tables)} ===")
    for ti, table in enumerate(doc.tables):
        print(f"\n=== TABLE {ti} ({len(table.rows)} rows x {len(table.columns)} cols) ===")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            print(" | ".join(cells))


if __name__ == "__main__":
    dump(sys.argv[1])
