#!/usr/bin/env python3
"""Extract text + formatting from 'Product Design - final.docx' to locate the
two big red section headings and the details under them."""
import docx
from docx.shared import RGBColor

DOC = "/home/z/my-project/upload/Product Design - final.docx"
d = docx.Document(DOC)

def run_color(r):
    try:
        c = r.font.color
        if c and c.rgb is not None:
            return str(c.rgb)
    except Exception:
        pass
    return None

print("=== PARAGRAPHS (style | size | bold | color | text) ===")
for i, p in enumerate(d.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    sizes, bolds, colors = set(), set(), set()
    for r in p.runs:
        if r.font.size:
            sizes.add(r.font.size.pt)
        if r.font.bold:
            bolds.add(True)
        col = run_color(r)
        if col:
            colors.add(col)
    style = p.style.name if p.style else "?"
    flag = ""
    if colors & {"FF0000", "C00000", "FF0100", "E36C0A"} or ("Heading" in style and colors):
        flag = "  <<< RED?"
    print(f"[{i:03d}] style={style!r} size={sorted(sizes)} bold={bool(bolds)} color={sorted(colors)} :: {txt[:150]}{flag}")

print("\n=== TABLES ===")
for ti, t in enumerate(d.tables):
    print(f"--- table {ti}: {len(t.rows)} rows x {len(t.columns)} cols ---")
    for ri, row in enumerate(t.rows[:20]):
        cells = [c.text.strip().replace("\n", " / ")[:60] for c in row.cells]
        print(f"  row{ri}: {cells}")
